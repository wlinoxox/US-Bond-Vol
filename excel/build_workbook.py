"""產生 us_bond_vol_rv_iv.xlsx：用 Bloomberg BDH 抓資料，計算 RV 與 IV − RV，拆解 10Y 為 TIPS 與 BEI，並對照 Kim-Wright TP。

用法：
    python excel/build_workbook.py            # 產生正式檔（BDH 公式，需在 Bloomberg Excel 開啟）
    python excel/build_workbook.py --sample   # 產生以模擬數據取代 BDH 的測試檔，用來驗證公式

Kim-Wright 資料來源：data/THREEFYTP10.csv（FRED 序列 THREEFYTP10，單位 %，日資料）。
DKW 資料來源：data/DKW_updates.csv（Fed staff D'Amico-Kim-Wei 拆解，單位 %，每月更新）。
更新任一資料時，覆蓋對應 CSV 後重新執行本腳本即可。DKW 為主軸，Kim-Wright 用於 DKW 尚未涵蓋的最新日期與交叉驗證。
"""

import csv
import math
import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
KW_CSV = ROOT / "data" / "THREEFYTP10.csv"
DKW_CSV = ROOT / "data" / "DKW_updates.csv"
DKW_START = date(2015, 1, 1)  # DKW 分頁只放這天以後的資料
DKW_COLS = ["exp.real.short.rate.10", "real.term.prem.10", "exp.inflation.10", "inflation.risk.prem.10",
            "tips.liq.prem.10", "nominal.yield.fitted.10"]
OUT = ROOT / "excel" / "us_bond_vol_rv_iv.xlsx"
SAMPLE_OUT = ROOT / "excel" / "us_bond_vol_rv_iv_sample.xlsx"  # 測試用，不納入版控

N_ROWS = int(__import__("os").environ.get("N_ROWS", 1500))  # Calc 預留的資料列數（約 5.7 年的週間日）
FIRST = 2  # 資料起始列
BBG_MAXR = 4000  # BBG_Data 每組 BDH 預留的列數（日曆日也夠用約 10 年）
# BBG_Data 每個 ticker 佔兩欄（日期、數值）；key 對應 Settings 回傳的參數名
BBG_PAIRS = [
    ("ust", "10Y 公債殖利率 (%)"), ("swap", "10Y SOFR Swap (%)"), ("move", "MOVE (bp)"),
    ("swpn", "3m10y Swaption IV (bp/y)"), ("tips", "10Y TIPS 實質殖利率 (%)"), ("bei", "10Y BEI (%)"),
    ("ust2", "2Y 公債殖利率 (%)"), ("ust5", "5Y 公債殖利率 (%)"), ("ust30", "30Y 公債殖利率 (%)"),
    ("hy", "HY OAS（BBG 原始單位）"), ("ig", "IG OAS（BBG 原始單位）"), ("cdxig", "CDX IG 5Y (bp)"),
    ("cdxhy", "CDX HY 5Y (bp)"),
]


def pair_cols(key):
    k = [p[0] for p in BBG_PAIRS].index(key)
    return get_column_letter(1 + 2 * k), get_column_letter(2 + 2 * k)
LAST = FIRST + N_ROWS - 1

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def load_kw():
    rows = []
    with open(KW_CSV, newline="") as f:
        for rec in csv.DictReader(f):
            v = rec["THREEFYTP10"].strip()
            if v and v != ".":
                rows.append((datetime.strptime(rec["observation_date"], "%Y-%m-%d").date(), float(v)))
    return rows


def load_dkw():
    rows = []
    with open(DKW_CSV, newline="") as f:
        lines = f.read().splitlines()
    start = next(i for i, l in enumerate(lines) if l.startswith('"date"'))
    for rec in csv.DictReader(lines[start:]):
        d = datetime.strptime(rec["date"], "%Y-%m-%d").date()
        if d < DKW_START:
            continue
        vals = [rec[c] for c in DKW_COLS]
        if any(v in ("", "NA") for v in vals[:4] + vals[5:]):
            continue
        rows.append((d, [None if v in ("", "NA") else float(v) for v in vals]))
    return rows


def build_readme(wb, kw_last, dkw_last):
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("美債 Vol 監控：RV、IV − RV 與 Kim-Wright TP", F_TITLE),
        ("", F_BASE),
        ("使用步驟", F_BOLD),
        ("1. 在已登入 Bloomberg 的電腦上用 Excel 開啟本檔（需 Bloomberg Excel Add-in）。", F_BASE),
        ("2. 到 Settings 分頁，檢查黃底藍字的輸入格：起訖日期、ticker、RV 視窗。", F_BASE),
        ("3. 3m10y Swaption Normal Vol 的 ticker 請到 VCUB <GO> 找 3M × 10Y（Normal / bp）那格，取得 ticker 後填入 Settings!B9。", F_BASE),
        ("4. 按 Bloomberg 工具列的 Refresh Workbook（或 Ctrl+Alt+Shift+R）讓 BBG_Data 的 BDH 抓資料，Calc 與 Dashboard 會自動計算。", F_BASE),
        ("", F_BASE),
        ("分頁說明", F_BOLD),
        ("Settings：所有輸入參數（黃底藍字＝可修改）。", F_BASE),
        ("BBG_Data：BDH 公式抓取的原始資料。每個 ticker 佔兩欄（日期、數值），Calc 依日期精確比對，不依列對齊。", F_BASE),
        ("KW_TP：Kim-Wright 10Y Term Premium（FRED：THREEFYTP10），使用者下載的靜態資料，最後一筆 " + kw_last.isoformat() + "。", F_BASE),
        ("DKW_10Y：D'Amico-Kim-Wei 10Y 拆解（Fed staff，每月更新），最後一筆 " + dkw_last.isoformat() + "。名目殖利率 = 預期實質短率 + 實質 TP + 預期通膨 + 通膨風險溢酬；名目 TP = 實質 TP + 通膨風險溢酬。", F_BASE),
        ("Calc：每日變動（bp）、滾動 RV、IV − RV、事後 VRP、Kim-Wright TP 對照、10Y = TIPS + BEI 的拆解，以及 DKW 各成分相對基準日的變動。", F_BASE),
        ("Dashboard：最新數值摘要與圖表（含曲線 2s10s / 5s30s 與信用利差）。", F_BASE),
        ("", F_BASE),
        ("計算定義", F_BOLD),
        ("每日變動 = (今日殖利率 − 前一日殖利率) × 100，單位 bp。", F_BASE),
        ("RV（年化 Normal Vol，bp/y）= 最近 N 個每日變動的樣本標準差 × √252。變動恰為 0（假日補值）不計入；視窗內有效筆數需達 N × Settings!B26。", F_BASE),
        ("IV − RV（即時近似）：Swaption IV − 10Y Swap 的 RV（63 日，對應 3 個月期限）；MOVE − 10Y 公債的 RV（21 日，對應 1 個月期限）。", F_BASE),
        ("事後 VRP：今天的 Swaption IV − 接下來 63 個交易日實際發生的 Swap RV。最近 63 天還沒有結果，會留空。", F_BASE),
        ("實質利率拆解（市場口徑）：10Y 名目 = 10Y TIPS 實質殖利率 + 10Y BEI。相對基準日的變動：Δ10Y = ΔTIPS（實質利率貢獻）+ ΔBEI（通膨預期貢獻）。", F_BASE),
        ("「名目 − TIPS − BEI」欄是檢查用，應接近 0；若差很多，代表 BEI ticker 的口徑和 10Y 名目 / TIPS 不一致。", F_BASE),
        ("注意：TIPS 殖利率含實質期限溢酬與流動性溢酬，BEI 含通膨風險溢酬；ΔTIPS 不等於「實質利率預期」的變動，要再拆需用 DKW 等模型。", F_BASE),
        ("", F_BASE),
        ("注意事項", F_BOLD),
        ("• Ticker 是依記憶填寫的預設值，第一次使用請先在 Terminal 確認（特別是 USOSFR10 Curncy、USGGT10Y Index、USGGBE10 Index 與 Swaption ticker）。", F_BASE),
        ("• 曾發生的問題：MOVE 的 BDH 即使指定 Days=W 仍回傳每個日曆日（含週末），依列對齊會錯位。現在每組 BDH 都輸出自己的日期並依日期比對，不受影響。", F_BASE),
        ("• HY / IG OAS：Bloomberg 回傳單位可能是 %（例如 3.18 = 318 bp），Calc 以 Settings!B25 的倍數換成 bp；若回傳已是 bp，把倍數改成 1。", F_BASE),
        ("• MOVE 是 1 個月期、2Y/5Y/10Y/30Y 加權的公債選擇權 IV，和 10Y 單一期限的 RV 比較只是近似。", F_BASE),
        ("• Calc 中無法計算的格子（RV 暖機期、預留的空白列）會顯示 N/A 錯誤值，這是刻意的，讓圖表在缺資料處斷開而不是掉到 0。", F_BASE),
        ("• 若在 Excel 365 看到公式變成 =@BDH(...) 且只抓到一筆，把 @ 刪掉重新輸入即可。", F_BASE),
        ("• TP 模型分工：DKW 為主軸（能拆實質 TP 與通膨風險溢酬）；Kim-Wright 為日資料，用於 DKW 尚未涵蓋的最新日期與交叉驗證。兩者都出自 Fed 的 Don Kim，並非完全獨立。", F_BASE),
        ("• DKW 資料不會自動更新：下載 DKW_updates.csv 覆蓋 data/ 後重新執行 build_workbook.py，或直接貼到 DKW_10Y 分頁。", F_BASE),
        ("• Kim-Wright 資料不會自動更新：到 FRED 下載新的 THREEFYTP10，貼到 KW_TP 分頁（日期在 A 欄、% 在 B 欄），C 欄公式往下複製即可。", F_BASE),
        ("• Calc 預留 " + str(N_ROWS) + " 列；若起始日期太早導致資料超過，請把 Calc 最後一列公式往下複製。", F_BASE),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font
        c.alignment = WRAP


def build_settings(wb, sample):
    ws = wb.create_sheet("Settings")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 70
    ws["A1"] = "參數設定"
    ws["A1"].font = F_TITLE
    header(ws, 2, ["項目", "值", "說明"])
    rows = [
        ("起始日期", date(2024, 1, 1), "BDH 起始日。RV 63 日需要約 3 個月暖機，建議比觀察期早至少 3 個月。"),
        ("結束日期", "=TODAY()", "預設今天，可改成固定日期。"),
        ("（空白）", None, None),
        ("10Y 公債殖利率 ticker", "USGG10YR Index", "Bloomberg 10Y 公債殖利率，單位 %。"),
        ("10Y SOFR Swap ticker", "USOSFR10 Curncy", "10Y SOFR OIS Swap，單位 %。請在 Terminal 確認 ticker。"),
        ("MOVE ticker", "MOVE Index", "ICE BofA MOVE Index，單位 bp。"),
        ("3m10y Swaption Normal Vol ticker", "SAMPLE" if sample else "", "請從 VCUB <GO> 取得 3M×10Y Normal Vol（bp/y）的 ticker 填入；留空則略過 Swaption 相關計算。"),
        ("（空白）", None, None),
        ("短期 RV 視窗（交易日）", 21, "約 1 個月，對應 MOVE（1 個月期選擇權）。"),
        ("長期 RV 視窗（交易日）", 63, "約 3 個月，對應 3m10y Swaption。"),
        ("年化天數", 252, "年化 = 日標準差 × √年化天數。"),
        ("10Y TIPS 實質殖利率 ticker", "USGGT10Y Index", "10Y TIPS 殖利率（實質利率，市場口徑），單位 %。請在 Terminal 確認 ticker。"),
        ("10Y BEI ticker", "USGGBE10 Index", "10Y Breakeven Inflation，單位 %。請在 Terminal 確認 ticker。"),
        ("分解基準日", date(2025, 12, 31), "計算「相對基準日的變動」的起點，例如前一年底或開戰日。"),
        ("基準日所在列（自動）", "=SUMPRODUCT(MAX(ISNUMBER(Calc!$A$2:$A$LASTROW)*(Calc!$A$2:$A$LASTROW<=$B$16)*ROW(Calc!$A$2:$A$LASTROW)))", "Calc 中日期 ≤ 基準日的最後一列；0 代表基準日早於資料起點。"),
        ("2Y 公債殖利率 ticker", "USGG2YR Index", "單位 %。"),
        ("5Y 公債殖利率 ticker", "USGG5YR Index", "單位 %。"),
        ("30Y 公債殖利率 ticker", "USGG30YR Index", "單位 %。"),
        ("HY OAS ticker", "LF98OAS Index", "Bloomberg US Corporate High Yield OAS。"),
        ("IG OAS ticker", "LUACOAS Index", "Bloomberg US Corporate IG OAS。"),
        ("CDX IG 5Y ticker", "IBOXUMAE CBBT Curncy", "單位 bp。"),
        ("CDX HY 5Y ticker", "IBOXHYSE CBBT Curncy", "單位 bp（spread）。"),
        ("OAS 換算倍數", 100, "BBG 回傳 % 時填 100（3.18 → 318 bp）；已是 bp 則填 1。"),
        ("RV 最少有效筆數比例", 0.85, "視窗內有效（非 0）變動至少要有視窗長度 × 此比例，否則顯示 N/A。"),
    ]
    for i, (label, value, note) in enumerate(rows, start=3):
        if label == "（空白）":
            continue
        ws.cell(row=i, column=1, value=label).font = F_BASE
        c = ws.cell(row=i, column=2, value=value)
        if isinstance(value, str) and value.startswith("="):
            c.font = F_BASE
        else:
            c.font = F_INPUT
            c.fill = FILL_INPUT
        if isinstance(value, str) and "LASTROW" in value:
            c.value = value.replace("LASTROW", str(LAST))
        if isinstance(value, date) or value == "=TODAY()":
            c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=3, value=note).font = F_BASE
        ws.cell(row=i, column=3).alignment = WRAP
    ws["B9"].comment = Comment("在 VCUB 中選 Normal Vol，找到 3M expiry × 10Y tenor 的格子，右鍵或滑鼠停留取得 ticker。", "README")
    ws["A28"] = "圖例：黃底藍字＝可修改的輸入；黑字＝公式；綠字＝引用其他分頁。"
    ws["A28"].font = F_BASE
    ws["A29"] = "範例：10Y 公債殖利率 ticker 填 USGG10YR Index；RV 視窗填整數，例如 21。"
    ws["A29"].font = F_BASE
    # 參數位置（供其他分頁引用）
    return {
        "start": "Settings!$B$3",
        "end": "Settings!$B$4",
        "ust": "Settings!$B$6",
        "swap": "Settings!$B$7",
        "move": "Settings!$B$8",
        "swpn": "Settings!$B$9",
        "w_short": "Settings!$B$11",
        "w_long": "Settings!$B$12",
        "ann": "Settings!$B$13",
        "tips": "Settings!$B$14",
        "bei": "Settings!$B$15",
        "base_row": "Settings!$B$17",
        "ust2": "Settings!$B$18",
        "ust5": "Settings!$B$19",
        "ust30": "Settings!$B$20",
        "hy": "Settings!$B$21",
        "ig": "Settings!$B$22",
        "cdxig": "Settings!$B$23",
        "cdxhy": "Settings!$B$24",
        "oas_mult": "Settings!$B$25",
        "minfrac": "Settings!$B$26",
    }


def build_bbg(wb, p, sample):
    ws = wb.create_sheet("BBG_Data")
    labels, widths = [], []
    for _, label in BBG_PAIRS:
        labels += [label.split(" (")[0].split("（")[0] + " 日期", label]
        widths += [12, 16]
    header(ws, 1, labels, widths)
    ws.row_dimensions[1].height = 45
    ovr = '"Dts=S","Days=W","Fill=P"'
    note_col = 2 * len(BBG_PAIRS) + 2
    if not sample:
        for key, _ in BBG_PAIRS:
            dc, _vc = pair_cols(key)
            f = f'BDH({p[key]},"PX_LAST",{p["start"]},{p["end"]},{ovr})'
            ws[f"{dc}2"] = f'=IF({p[key]}="","",{f})' if key == "swpn" else f"={f}"
            ws[f"{dc}2"].font = F_LINK
        notes = [
            "每組 BDH 輸出兩欄：日期（Dts=S）與數值；Calc 用日期精確比對（MATCH），不依列對齊。",
            "原因：MOVE 等指數即使指定 Days=W，Bloomberg 仍可能回傳每個日曆日，依列對齊會錯位。",
            f"每組預留 {BBG_MAXR - 1} 列；請勿在公式下方輸入任何內容，以免擋住 BDH 輸出。",
        ]
        ws.cell(row=1, column=note_col, value="說明").font = F_BOLD
        for i, n in enumerate(notes, start=2):
            ws.cell(row=i, column=note_col, value=n).font = F_BASE
        ws.column_dimensions[get_column_letter(note_col)].width = 90
    else:
        # 測試檔：用模擬數據取代 BDH。MOVE 刻意放成「每個日曆日」、HY OAS 刻意缺幾天，用來驗證依日期比對。
        rng = random.Random(42)
        days = [date(2024, 1, 1) + timedelta(days=i) for i in range((date(2026, 10, 2) - date(2024, 1, 1)).days + 1)]
        wk = [d for d in days if d.weekday() < 5]
        lvl = {"ust": 3.9, "swap": 3.6, "move": 110.0, "swpn": 95.0, "bei": 2.3, "ust2": 4.2, "ust5": 3.8,
               "ust30": 4.1, "hy": 3.4, "ig": 0.95, "cdxig": 60.0, "cdxhy": 350.0}
        data = {k: [] for k, _ in BBG_PAIRS}
        move_by_day = {}
        for d in wk:
            shock = rng.gauss(0, 0.055)
            lvl["ust"] += shock
            lvl["swap"] += shock * 0.95 + rng.gauss(0, 0.01)
            lvl["bei"] += rng.gauss(0, 0.02)
            lvl["ust2"] += shock * 0.8 + rng.gauss(0, 0.02)
            lvl["ust5"] += shock * 0.9 + rng.gauss(0, 0.015)
            lvl["ust30"] += shock * 1.05 + rng.gauss(0, 0.015)
            lvl["move"] = max(60, lvl["move"] + rng.gauss(0, 2.5))
            lvl["swpn"] = max(50, lvl["swpn"] + rng.gauss(0, 1.8))
            lvl["hy"] = max(2.0, lvl["hy"] + rng.gauss(0, 0.04))
            lvl["ig"] = max(0.5, lvl["ig"] + rng.gauss(0, 0.01))
            lvl["cdxig"] = max(40, lvl["cdxig"] + rng.gauss(0, 0.8))
            lvl["cdxhy"] = max(250, lvl["cdxhy"] + rng.gauss(0, 5))
            for k in ("ust", "swap", "swpn", "ust2", "ust5", "ust30", "ig", "cdxig", "cdxhy"):
                data[k].append((d, round(lvl[k], 4)))
            data["tips"].append((d, round(lvl["ust"] - lvl["bei"], 4)))
            data["bei"].append((d, round(lvl["bei"], 4)))
            if rng.random() > 0.03:
                data["hy"].append((d, round(lvl["hy"], 4)))
            move_by_day[d] = round(lvl["move"], 2)
        last = None
        for d in days:  # MOVE：每個日曆日，週末沿用前值
            last = move_by_day.get(d, last)
            if last is not None:
                data["move"].append((d, last))
        for key, _ in BBG_PAIRS:
            dc, vc = pair_cols(key)
            for i, (d, v) in enumerate(data[key], start=FIRST):
                c = ws[f"{dc}{i}"]
                c.value, c.number_format, c.font = d, "yyyy-mm-dd", F_INPUT
                ws[f"{vc}{i}"].value = v
                ws[f"{vc}{i}"].font = F_INPUT
    ws.freeze_panes = "A2"


def build_kw(wb, kw):
    ws = wb.create_sheet("KW_TP")
    header(ws, 1, ["日期", "Kim-Wright 10Y TP (%)", "Kim-Wright 10Y TP (bp)"], [12, 22, 22])
    for i, (d, v) in enumerate(kw, start=2):
        ws.cell(row=i, column=1, value=d).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=2, value=v).font = F_INPUT
        ws.cell(row=i, column=3, value=f"=B{i}*100").number_format = "0.0"
        ws.cell(row=i, column=1).font = F_INPUT
    ws["E1"] = "來源：FRED 序列 THREEFYTP10（Kim-Wright 三因子模型 10Y Term Premium），由使用者下載；已刪除空值（假日）。"
    ws["E1"].font = F_BASE
    ws.freeze_panes = "A2"
    return len(kw) + 1


def build_dkw(wb, dkw):
    ws = wb.create_sheet("DKW_10Y")
    header(ws, 1, ["日期", "預期實質短率 (%)", "實質 TP (%)", "預期通膨 (%)", "通膨風險溢酬 (%)",
                   "TIPS 流動性溢酬 (%)", "名目殖利率 fitted (%)", "名目 TP (%)"],
           [12, 16, 12, 12, 16, 18, 18, 12])
    for i, (d, vals) in enumerate(dkw, start=2):
        c = ws.cell(row=i, column=1, value=d)
        c.number_format = "yyyy-mm-dd"
        c.font = F_INPUT
        for j, v in enumerate(vals, start=2):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F_INPUT
            c.number_format = "0.0000"
        c = ws.cell(row=i, column=8, value=f"=C{i}+E{i}")
        c.number_format = "0.0000"
    ws["J1"] = ("來源：Fed staff DKW_updates.csv（D'Amico, Kim & Wei 2018；Kim, Walsh & Wei 2019 FEDS Notes），"
                f"由使用者下載；只保留 {DKW_START.isoformat()} 之後且四個成分齊全的日期。非 Fed 官方統計，可能修正。")
    ws["J1"].font = F_BASE
    ws.freeze_panes = "A2"
    return len(dkw) + 1


def build_calc(wb, p, kw_last_row, dkw_last_row):
    ws = wb.create_sheet("Calc")
    labels = [
        "日期", "10Y 公債 (%)", "10Y Swap (%)", "MOVE (bp)", "Swaption IV (bp/y)",
        "Δ 10Y 公債 (bp)", "Δ 10Y Swap (bp)",
        "RV 公債 短期 (bp/y)", "RV 公債 長期 (bp/y)", "RV Swap 短期 (bp/y)", "RV Swap 長期 (bp/y)",
        "MOVE − RV 公債短期", "Swaption IV − RV Swap 長期",
        "未來長期視窗 Swap RV (bp/y)", "事後 VRP：Swaption IV − 未來 RV",
        "Kim-Wright TP (bp)",
        "10Y TIPS (%)", "10Y BEI (%)", "檢查：名目 − TIPS − BEI (bp)",
        "Δ10Y 名目 相對基準日 (bp)", "ΔTIPS 實質利率貢獻 (bp)", "ΔBEI 通膨預期貢獻 (bp)", "ΔKim-Wright TP 相對基準日 (bp)",
        "DKW Δ預期實質短率 (bp)", "DKW Δ實質 TP (bp)", "DKW Δ預期通膨 (bp)", "DKW Δ通膨風險溢酬 (bp)", "DKW Δ名目 TP (bp)",
        "2Y 公債 (%)", "5Y 公債 (%)", "30Y 公債 (%)", "2s10s (bp)", "5s30s (bp)",
        "HY OAS (bp)", "IG OAS (bp)", "CDX IG 5Y (bp)", "CDX HY 5Y (bp)",
    ]
    header(ws, 1, labels, [12, 11, 11, 10, 12, 11, 11, 12, 12, 12, 12, 13, 14, 14, 15, 12, 11, 11, 13, 13, 13, 13, 14,
                           13, 12, 12, 13, 12, 11, 11, 11, 11, 11, 11, 11, 11, 11])
    ws.row_dimensions[1].height = 45
    kw_a = f"KW_TP!$A$2:$A${kw_last_row}"
    kw_c = f"KW_TP!$C$2:$C${kw_last_row}"
    ws_, wl, ann = p["w_short"], p["w_long"], p["ann"]

    mf = p["minfrac"]

    def rv(col, r, win):
        rng_ = f"OFFSET({col}{r},1-{win},0,{win},1)"
        return (f'=IF(OR(NOT(ISNUMBER($A{r})),ROW()-{FIRST}<{win}),NA(),'
                f'IF(COUNT({rng_})<ROUNDUP({win}*{mf},0),NA(),STDEV({rng_})*SQRT({ann})))')

    def look(key, r, mult=None):
        dc, vc = pair_cols(key)
        val = (f"INDEX(BBG_Data!${vc}$2:${vc}${BBG_MAXR},"
               f"MATCH($A{r},BBG_Data!${dc}$2:${dc}${BBG_MAXR},0))")
        val = f"{val}*{mult}" if mult else f"1*{val}"
        return f"=IF(NOT(ISNUMBER($A{r})),NA(),IFERROR({val},NA()))"

    for r in range(FIRST, LAST + 1):
        ws[f"A{r}"] = f'=IF(ISNUMBER(BBG_Data!A{r}),BBG_Data!A{r},"")'
        ws[f"B{r}"] = f'=IF(AND(ISNUMBER($A{r}),ISNUMBER(BBG_Data!B{r})),BBG_Data!B{r},NA())'
        for col, key in (("C", "swap"), ("D", "move"), ("E", "swpn")):
            ws[f"{col}{r}"] = look(key, r)
        for src, dst in (("B", "F"), ("C", "G")):
            if r == FIRST:
                ws[f"{dst}{r}"] = '=""'
            else:
                ws[f"{dst}{r}"] = (f'=IF(AND(ISNUMBER({src}{r}),ISNUMBER({src}{r-1})),'
                                   f'IF({src}{r}={src}{r-1},"",({src}{r}-{src}{r-1})*100),"")')
        ws[f"H{r}"] = rv("F", r, ws_)
        ws[f"I{r}"] = rv("F", r, wl)
        ws[f"J{r}"] = rv("G", r, ws_)
        ws[f"K{r}"] = rv("G", r, wl)
        ws[f"L{r}"] = f"=IF(AND(ISNUMBER(D{r}),ISNUMBER(H{r})),D{r}-H{r},NA())"
        ws[f"M{r}"] = f"=IF(AND(ISNUMBER(E{r}),ISNUMBER(K{r})),E{r}-K{r},NA())"
        fwd = f"OFFSET(G{r},1,0,{wl},1)"
        ws[f"N{r}"] = (f"=IF(NOT(ISNUMBER($A{r})),NA(),IF(COUNT({fwd})<ROUNDUP({wl}*{mf},0),NA(),"
                       f"STDEV({fwd})*SQRT({ann})))")
        ws[f"O{r}"] = f"=IF(AND(ISNUMBER(E{r}),ISNUMBER(N{r})),E{r}-N{r},NA())"
        ws[f"P{r}"] = (f'=IF(NOT(ISNUMBER($A{r})),NA(),IF(OR($A{r}<MIN({kw_a}),$A{r}>MAX({kw_a})),NA(),'
                       f'INDEX({kw_c},MATCH($A{r},{kw_a},1))))')
        for col, key in (("Q", "tips"), ("R", "bei")):
            ws[f"{col}{r}"] = look(key, r)
        for col, key in (("AC", "ust2"), ("AD", "ust5"), ("AE", "ust30")):
            ws[f"{col}{r}"] = look(key, r)
            ws[f"{col}{r}"].number_format = "0.000"
            ws[f"{col}{r}"].font = F_LINK
        ws[f"AF{r}"] = f"=IF(AND(ISNUMBER(B{r}),ISNUMBER(AC{r})),(B{r}-AC{r})*100,NA())"
        ws[f"AG{r}"] = f"=IF(AND(ISNUMBER(AE{r}),ISNUMBER(AD{r})),(AE{r}-AD{r})*100,NA())"
        for col, key, mult in (("AH", "hy", p["oas_mult"]), ("AI", "ig", p["oas_mult"]),
                               ("AJ", "cdxig", None), ("AK", "cdxhy", None)):
            ws[f"{col}{r}"] = look(key, r, mult)
            ws[f"{col}{r}"].font = F_LINK
        for col in ("AF", "AG", "AH", "AI", "AJ", "AK"):
            ws[f"{col}{r}"].number_format = "0.0"
        for col in ("AF", "AG"):
            ws[f"{col}{r}"].font = F_BASE
        ws[f"S{r}"] = f"=IF(AND(ISNUMBER(B{r}),ISNUMBER(Q{r}),ISNUMBER(R{r})),(B{r}-Q{r}-R{r})*100,NA())"
        br = p["base_row"]
        for dst, src in (("T", "B"), ("U", "Q"), ("V", "R")):
            ws[f"{dst}{r}"] = (f"=IF(AND(ISNUMBER({src}{r}),{br}>0),IF(ISNUMBER(INDEX(${src}:${src},{br})),"
                               f"({src}{r}-INDEX(${src}:${src},{br}))*100,NA()),NA())")
        ws[f"W{r}"] = (f"=IF(AND(ISNUMBER(P{r}),{br}>0),IF(ISNUMBER(INDEX($P:$P,{br})),"
                       f"P{r}-INDEX($P:$P,{br}),NA()),NA())")
        dkw_a = f"DKW_10Y!$A$2:$A${dkw_last_row}"
        for dst, src in (("X", "B"), ("Y", "C"), ("Z", "D"), ("AA", "E"), ("AB", "H")):
            rng_ = f"DKW_10Y!${src}$2:${src}${dkw_last_row}"
            ws[f"{dst}{r}"] = (f"=IF(NOT(ISNUMBER($A{r})),NA(),IF(OR($A{r}<MIN({dkw_a}),$A{r}>MAX({dkw_a}),"
                               f"Settings!$B$16<MIN({dkw_a})),NA(),"
                               f"(INDEX({rng_},MATCH($A{r},{dkw_a},1))-INDEX({rng_},MATCH(Settings!$B$16,{dkw_a},1)))*100))")
            ws[f"{dst}{r}"].number_format = "0.0"
            ws[f"{dst}{r}"].font = F_LINK
        ws[f"A{r}"].number_format = "yyyy-mm-dd"
        for col in "QR":
            ws[f"{col}{r}"].number_format = "0.000"
            ws[f"{col}{r}"].font = F_LINK
        for col in "STUVW":
            ws[f"{col}{r}"].number_format = "0.0"
            ws[f"{col}{r}"].font = F_BASE
        for col in "BC":
            ws[f"{col}{r}"].number_format = "0.000"
        for col in "DEFGHIJKLMNOP":
            ws[f"{col}{r}"].number_format = "0.0"
        for col in "ABCDE":
            ws[f"{col}{r}"].font = F_LINK
        for col in "FGHIJKLMNO":
            ws[f"{col}{r}"].font = F_BASE
        ws[f"P{r}"].font = F_LINK
    ws.freeze_panes = "B2"


def build_dashboard(wb, kw_last_row, dkw_last_row):
    ws = wb.create_sheet("Dashboard")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60
    ws["A1"] = "最新數值摘要"
    ws["A1"].font = F_TITLE
    header(ws, 2, ["指標", "數值", "說明"])
    calc = f"Calc!$A${FIRST}:$A${LAST}"
    kw_a = f"KW_TP!$A$2:$A${kw_last_row}"
    kw_c = f"KW_TP!$C$2:$C${kw_last_row}"

    dkw_a = f"DKW_10Y!$A$2:$A${dkw_last_row}"

    def dkw_d(col):
        rng_ = f"DKW_10Y!${col}$2:${col}${dkw_last_row}"
        return (f'=IFERROR((INDEX({rng_},COUNT({dkw_a}))-INDEX({rng_},MATCH(Settings!$B$16,{dkw_a},1)))*100,"n/a")')

    def at(col):
        return f"=IFERROR(INDEX(Calc!${col}${FIRST}:${col}${LAST},$B$3),\"n/a\")"

    def since(col):
        return (f"=IFERROR(INDEX(Calc!${col}${FIRST}:${col}${LAST},$B$3)"
                f"-INDEX(Calc!${col}:${col},Settings!$B$17),\"n/a\")")

    rows = [
        ("最新資料列（Calc 內序號）", f"=MATCH(9.99E+307,{calc})", "Calc 中最後一個有日期的列。", "0"),
        ("最新日期", f"=INDEX({calc},B3)", "", "yyyy-mm-dd"),
        ("10Y 公債殖利率 (%)", at("B"), "", "0.000"),
        ("MOVE (bp)", at("D"), "", "0.0"),
        ("RV 公債 短期 (bp/y)", at("H"), "對照 MOVE。", "0.0"),
        ("MOVE − RV 公債短期", at("L"), "正值＝選擇權比實際波動貴。", "0.0"),
        ("Swaption IV (bp/y)", at("E"), "需先在 Settings!B9 填 ticker。", "0.0"),
        ("RV Swap 長期 (bp/y)", at("K"), "對照 3m10y Swaption。", "0.0"),
        ("Swaption IV − RV Swap 長期", at("M"), "即時版 VRP 近似。", "0.0"),
        ("Kim-Wright TP 最新 (bp)", f"=INDEX({kw_c},COUNT({kw_a}))", "KW_TP 分頁最後一筆。", "0.0"),
        ("Kim-Wright TP 最新日期", f"=INDEX({kw_a},COUNT({kw_a}))", "", "yyyy-mm-dd"),
        ("Kim-Wright TP 今年以來變動 (bp)",
         f"=B12-INDEX({kw_c},MATCH(DATE(YEAR(B13)-1,12,31),{kw_a},1))", "相對前一年最後一筆。", "0.0"),
        ("10Y TIPS 實質殖利率 (%)", at("Q"), "市場口徑的實質利率。", "0.000"),
        ("10Y BEI (%)", at("R"), "", "0.000"),
        ("分解基準日", "=Settings!B16", "在 Settings!B16 修改。", "yyyy-mm-dd"),
        ("Δ10Y 名目（基準日以來，bp）", at("T"), "", "0.0"),
        ("ΔTIPS：實質利率貢獻 (bp)", at("U"), "含實質期限溢酬與流動性溢酬。", "0.0"),
        ("ΔBEI：通膨預期貢獻 (bp)", at("V"), "含通膨風險溢酬。", "0.0"),
        ("實質利率貢獻占比", '=IFERROR(B19/B18,"n/a")', "ΔTIPS ÷ Δ10Y。", "0%"),
        ("ΔKim-Wright TP（基準日以來，bp）", at("W"), "模型口徑，和上面的市場口徑不能直接相加；Kim-Wright 資料較晚時顯示 n/a。", "0.0"),
        ("DKW 最新日期", f"=INDEX({dkw_a},COUNT({dkw_a}))", "DKW 每月更新；以下為基準日到此日的變動。", "yyyy-mm-dd"),
        ("DKW Δ名目殖利率 fitted (bp)", dkw_d("G"), "= 下面四個成分的加總。", "0.0"),
        ("DKW Δ預期實質短率 (bp)", dkw_d("B"), "政策路徑與 r* 的預期。", "0.0"),
        ("DKW Δ實質 TP (bp)", dkw_d("C"), "TIPS「實質利率」上升中屬於期限溢酬的部分。", "0.0"),
        ("DKW Δ預期通膨 (bp)", dkw_d("D"), "", "0.0"),
        ("DKW Δ通膨風險溢酬 (bp)", dkw_d("E"), "", "0.0"),
        ("DKW Δ名目 TP (bp)", dkw_d("H"), "= 實質 TP + 通膨風險溢酬，可與 Kim-Wright 對照。", "0.0"),
        ("DKW 名目 TP 占名目變動比例", '=IFERROR(B29/B24,"n/a")', "", "0%"),
        ("DKW 實質 TP 占實質部分比例", '=IFERROR(B26/(B25+B26),"n/a")', "實質部分 = 預期實質短率 + 實質 TP。", "0%"),
        ("2s10s (bp)", at("AF"), "正值愈大 = 曲線愈陡。", "0.0"),
        ("2s10s 基準日以來變動 (bp)", since("AF"), "正值 = 變陡。", "0.0"),
        ("5s30s (bp)", at("AG"), "", "0.0"),
        ("5s30s 基準日以來變動 (bp)", since("AG"), "", "0.0"),
        ("HY OAS (bp)", at("AH"), "", "0.0"),
        ("HY OAS 基準日以來變動 (bp)", since("AH"), "", "0.0"),
        ("IG OAS (bp)", at("AI"), "", "0.0"),
        ("CDX IG 5Y (bp)", at("AJ"), "", "0.0"),
        ("CDX HY 5Y (bp)", at("AK"), "", "0.0"),
    ]
    for i, (label, formula, note, fmt) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label).font = F_BASE
        c = ws.cell(row=i, column=2, value=formula)
        c.font = F_LINK
        c.number_format = fmt
        ws.cell(row=i, column=3, value=note).font = F_BASE

    def line(title, cols, anchor, y_title):
        ch = LineChart()
        ch.title = title
        ch.height, ch.width = 7.5, 16
        ch.y_axis.title = y_title
        ch.x_axis = DateAxis(crossAx=100)
        ch.x_axis.number_format = "yyyy-mm"
        ch.x_axis.majorTimeUnit = "months"
        calc_ws = wb["Calc"]
        for col in cols:
            ref = Reference(calc_ws, min_col=col, min_row=1, max_row=LAST)
            ch.add_data(ref, titles_from_data=True)
        ch.set_categories(Reference(calc_ws, min_col=1, min_row=FIRST, max_row=LAST))
        for s in ch.series:
            s.smooth = False
            s.graphicalProperties.line.width = 15000
        ws.add_chart(ch, anchor)

    line("3m10y Swaption IV vs. 10Y Swap RV", [5, 10, 11], "E2", "bp/y")
    line("MOVE vs. 10Y 公債 RV（短期）", [4, 8], "E18", "bp")
    line("IV − RV", [12, 13], "E34", "bp/y")
    line("事後 VRP 與 Kim-Wright TP", [15, 16], "E50", "bp")
    line("10Y 拆解：相對基準日的變動", [20, 21, 22, 23], "E66", "bp")
    line("DKW 10Y 拆解：相對基準日的變動", [24, 25, 26, 27, 28], "E82", "bp")
    line("曲線：2s10s 與 5s30s", [32, 33], "E98", "bp")
    line("信用利差：HY / IG OAS", [34, 35], "E114", "bp")


def main():
    sample = "--sample" in sys.argv
    kw = load_kw()
    dkw = load_dkw()
    wb = Workbook()
    build_readme(wb, kw[-1][0], dkw[-1][0])
    p = build_settings(wb, sample)
    build_bbg(wb, p, sample)
    kw_last_row = build_kw(wb, kw)
    dkw_last_row = build_dkw(wb, dkw)
    build_calc(wb, p, kw_last_row, dkw_last_row)
    build_dashboard(wb, kw_last_row, dkw_last_row)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, size=c.font.size or 10, bold=c.font.bold, color=c.font.color)
    wb.calculation.fullCalcOnLoad = True
    out = SAMPLE_OUT if sample else OUT
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
